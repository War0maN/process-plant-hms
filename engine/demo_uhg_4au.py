"""
SP зөвлөмжийн БҮРЭН жишээ — Ухаа Худаг 4AU бодит дээж (Book2) дээр.

Энэ нь engine-ийг баталгаажсан config-той, бодит лабын washability дээр
ажиллуулж, оператор юу харахыг үзүүлнэ. Дахин ажиллуулах:  python -m engine.demo_uhg_4au
"""
from pathlib import Path

from openpyxl import load_workbook

from engine import washability as w
from engine.washability import Fraction

BOOK2 = Path(__file__).resolve().parent.parent / "docs" / "refs" / "Book2.xlsx"

# --- Баталгаажсан config (reference_plant_a.yaml-аас) ---
TARGET_COKING = 10.5      # зорилтот баяжмалын үнс
CONTRACT_BASE = 10.0      # гэрээний үндсэн үнс
CONTRACT_PENALTY = 10.5   # >10.5% → $1.5/т торгууль
CONTRACT_REJECT = 11.5    # >11.5% → татгалзах
TARGET_THERMAL = 17.0     # эрчим хүчний дунд бүтээгдэхүүний үнс
SP_PRIMARY = 1.38         # одоогийн ажлын орчны нягт (gauge typical)
SP_SECONDARY = 1.50
NGM_HW = 0.05             # баталгаажсан конвенц


def load_4au():
    wb = load_workbook(BOOK2, data_only=True)
    ws = wb["Float sink"]
    fr, prev = [], 1.25
    for r in range(11, 25):
        hi = ws[f"AK{r}"].value
        mass = ws[f"AL{r}"].value
        ash = ws[f"AM{r}"].value
        if hi is None or mass is None:
            continue
        fr.append(Fraction(lo=prev, hi=hi, mass=mass, ash=ash))
        prev = hi
    return fr


def line(c="─", n=66):
    print(c * n)


def main():
    fr = load_4au()
    total_ash = w.weighted_ash(fr)

    line("═")
    print(" SP ЗӨВЛӨМЖ — Ухаа Худаг 4AU дээж (бодит washability, Book2)")
    line("═")
    print(f" Тэжээлийн нийт үнс (washability):  {total_ash:.2f}%")
    print(f" Зорилтот коксжих үнс:              {TARGET_COKING:.1f}%")
    print(f" Гэрээ: үндсэн {CONTRACT_BASE}% | торгууль >{CONTRACT_PENALTY}% | "
          f"татгалзах >{CONTRACT_REJECT}%")
    print()

    # --- A. Гарц-Үнс хүснэгт (хэд хэдэн cut дээр) ---
    line()
    print(" A. Гарц–Үнс–NGM (cut density тус бүрээр)")
    line()
    print(f" {'Cut RD':>7} {'Гарц %':>8} {'Үнс %':>8} {'NGM±05 %':>9}  Тэмдэглэл")
    for d in [1.35, 1.375, 1.39, 1.40, 1.425, 1.45]:
        y = w.yield_at(fr, d)
        a = w.clean_ash_at(fr, d)
        ngm = w.ngm_at(fr, d, NGM_HW)
        note = ""
        if a > CONTRACT_REJECT:
            note = "✗ татгалзах"
        elif a > CONTRACT_PENALTY:
            note = "⚠ торгууль"
        elif a > CONTRACT_BASE:
            note = "• base дээгүүр"
        print(f" {d:>7.3f} {y:>8.2f} {a:>8.2f} {ngm:>9.1f}  {note}")
    print()

    # --- B. Зорилтот 10.5%-д хүрэх cut ---
    line()
    print(" B. Зорилтот коксжих үнс 10.5% → шаардлагатай cut")
    line()
    cut105 = w.cut_for_product_ash(fr, TARGET_COKING)
    cut100 = w.cut_for_product_ash(fr, CONTRACT_BASE)
    y105, y100 = w.yield_at(fr, cut105), w.yield_at(fr, cut100)
    ngm105 = w.ngm_at(fr, cut105, NGM_HW)
    print(f" Үнс 10.5% (penalty ирмэг): cut={cut105:.3f} → гарц {y105:.2f}%, NGM {ngm105:.1f}%")
    print(f" Үнс 10.0% (аюулгүй base):  cut={cut100:.3f} → гарц {y100:.2f}%")
    print(f" → 10.5 vs 10.0 дээр гарцын зөрүү: {y105 - y100:+.2f}% "
          f"(0.5% үнсний төлөө {y105 - y100:.1f}% гарц)")
    print()

    # --- C. Бүрэн SP зөвлөмж (хэмжсэн үнсгүй — washability таамаг) ---
    line()
    print(" C. SP зөвлөмж — тэжээл washability-тэй ойролцоо (online ash алга)")
    line()
    rec = w.recommend([{"code": "4AU", "ratio": 100}], {"4AU": fr},
                      target_coking=TARGET_COKING, target_thermal=TARGET_THERMAL,
                      sp_primary=SP_PRIMARY, sp_secondary=SP_SECONDARY,
                      measured_ash=None, ngm_window=NGM_HW)
    p, s, b = rec["primary"], rec["secondary"], rec["balance"]
    print(f" КОКСЖИХ:  cut={p['corr_cut']:.3f}  орчны SP≈{p['reco_sp']:.3f}  "
          f"гарц={b['coking']['yield']:.1f}%  үнс={b['coking']['ash']:.2f}%  NGM={p['ngm']:.1f}%")
    print(f" ЭРЧИМ ХҮЧ: cut={s['corr_cut']:.3f}  орчны SP≈{s['reco_sp']:.3f}  "
          f"гарц={b['thermal']['yield']:.1f}%  үнс={b['thermal']['ash']:.2f}%")
    print(f" ХАЯГДАЛ:                              "
          f"гарц={b['reject']['yield']:.1f}%  үнс={b['reject']['ash']:.2f}%")
    print()

    # --- D. Online ash сценари: тэжээл бохирдвол ---
    line()
    print(" D. Online ash (AIT_105) тэжээл бохирдлыг мэдрэв: 24.76% → 27.0%")
    line()
    rec2 = w.recommend([{"code": "4AU", "ratio": 100}], {"4AU": fr},
                       target_coking=TARGET_COKING, target_thermal=TARGET_THERMAL,
                       sp_primary=SP_PRIMARY, sp_secondary=SP_SECONDARY,
                       measured_ash=27.0, ngm_window=NGM_HW)
    p2 = rec2["primary"]
    print(f" Tilt k={rec2['tilt_k']:+.2f} → корр. cut={p2['corr_cut']:.3f} "
          f"(base {p2['base_cut']:.3f}, nudge {p2['nudge']:+.3f})")
    print(f" → Орчны SP зөвлөмж: {SP_PRIMARY:.3f} → {p2['reco_sp']:.3f} "
          f"({p2['nudge']:+.3f})  [чанараа барихын тулд нягт буулгана]")
    print()

    # --- E. NGM хүндрэлийн дүгнэлт ---
    line()
    print(" E. Угаагдах чанарын дүгнэлт")
    line()
    print(f" Ажлын cut {cut105:.3f} дээр NGM(±0.05)={ngm105:.1f}% — энэ нь МАШ ӨНДӨР.")
    print(" 4AU бол туйлын хэцүү угаагдах нүүрс: cut орчимд асар их хагас-чулуулаг.")
    print(" → Орчны нягтыг маш хатуу (±0.005) барих ёстой; бага савлахад чулуу үсэрнэ.")
    line("═")


if __name__ == "__main__":
    main()
