"""
DMC Inspection log ачаалагч — циклоны амсрын элэгдлийн сүүлийн хэмжилт.

Ухаа Худагийн 'Primary/Secondary DMC Inspection' xlsx: мөр бүр нэг үзлэг —
Date, Feed on Cyclone (тонн), Inspection bore (Inlet/Vortex/Nozzle(spigot)/Shroud, мм).
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional


def _max_bore(v) -> Optional[float]:
    """'590', '575/565' зэргээс хамгийн ИХ (хамгийн элэгдсэн) утгыг авах."""
    if v is None:
        return None
    s = str(v).replace(",", ".")
    nums = []
    for part in s.replace("\\", "/").split("/"):
        try:
            nums.append(float(part.strip()))
        except ValueError:
            pass
    return max(nums) if nums else None


def latest_dmc_wear(path, sheet=None, cols=None) -> dict:
    """
    Хамгийн сүүлийн (доод талын) bore хэмжилт бүхий мөрийг буцаах.

    cols: {"vortex":6,"spigot":7,"shroud":8,"feed":4,"date":2} (1-based).
    Default нь Primary sheet-ийн байрлал.
    """
    from openpyxl import load_workbook

    cols = cols or {"date": 2, "feed": 4, "vortex": 6, "spigot": 7, "shroud": 8}
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[1]
    latest = {"date": None, "feed_tonnage": None, "vortex": None,
              "spigot": None, "shroud": None}
    for r in range(1, ws.max_row + 1):
        v = _max_bore(ws.cell(r, cols["vortex"]).value)
        g = _max_bore(ws.cell(r, cols["spigot"]).value)
        h = _max_bore(ws.cell(r, cols["shroud"]).value)
        if v is None and g is None and h is None:
            continue
        latest = {
            "date": ws.cell(r, cols["date"]).value,
            "feed_tonnage": ws.cell(r, cols["feed"]).value,
            "vortex": v, "spigot": g, "shroud": h,
        }
    return latest
